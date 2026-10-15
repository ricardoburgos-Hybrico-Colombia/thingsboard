# -*- coding: utf-8 -*-
r"""
Generador del Libro Excel Exclusivo para la Matriz de Alarmas de los 15 Sitios Telecom
Archivo destino: C:\Users\PC\Documents\AI_Engineer\20.LEER_DATA_API_ENERCLO\Matriz_Alarmas_Sitios_Telecom.xlsx
También restaura la hoja de telecom en Mapeo_Telemetria_Variables_y_Formulas_GAORI.xlsx.
"""

import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

DOC_DIR = r"C:\Users\PC\Documents\AI_Engineer\20.LEER_DATA_API_ENERCLO"
NEW_EXCEL_PATH = os.path.join(DOC_DIR, "Matriz_Alarmas_Sitios_Telecom.xlsx")
GAORI_EXCEL_PATH = os.path.join(DOC_DIR, "Mapeo_Telemetria_Variables_y_Formulas_GAORI.xlsx")

from telecom_data import SITES_15_LIST, TELECOM_ALARMS

def generar_excel_dedicado_telecom():
    print(f"[1/3] Creando libro Excel dedicado para Telecom en: {NEW_EXCEL_PATH}...")
    wb = openpyxl.Workbook()
    default_sheet = wb.active

    # Paleta Corporativa HyVision Telecom (Verde Oscuro / Esmeralda Industrial)
    fill_dark    = PatternFill(start_color="0B1A10", end_color="0B1A10", fill_type="solid")
    fill_header  = PatternFill(start_color="1B4332", end_color="1B4332", fill_type="solid")
    fill_sub     = PatternFill(start_color="2D6A4F", end_color="2D6A4F", fill_type="solid")
    fill_zebra   = PatternFill(start_color="F4FBF7", end_color="F4FBF7", fill_type="solid")
    fill_white   = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    fill_crit    = PatternFill(start_color="FDE8E8", end_color="FDE8E8", fill_type="solid")
    fill_warn    = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")

    font_title   = Font(name="Segoe UI", size=15, bold=True, color="FFFFFF")
    font_subtitle= Font(name="Segoe UI", size=10, italic=True, color="E2E8F0")
    font_head    = Font(name="Segoe UI", size=10.5, bold=True, color="FFFFFF")
    font_body    = Font(name="Segoe UI", size=9.5, color="1E293B")
    font_bold    = Font(name="Segoe UI", size=9.5, bold=True, color="0F172A")
    font_code    = Font(name="Consolas", size=9.5, bold=True, color="0369A1")
    font_cond    = Font(name="Segoe UI", size=9.5, color="1E3A8A")
    font_crit    = Font(name="Segoe UI", size=9.5, bold=True, color="991B1B")
    font_warn    = Font(name="Segoe UI", size=9.5, bold=True, color="92400E")

    border_thin  = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    # =========================================================================
    # HOJA 1: Matriz_Alarmas_Telecom (26 ALARMAS COMPLETAS)
    # =========================================================================
    ws1 = wb.create_sheet(title="Matriz_Alarmas_Telecom")
    ws1.views.sheetView[0].showGridLines = True

    # Banner superior
    ws1.merge_cells("A1:K1")
    ws1["A1"] = f"MATRIZ SCADA DE {len(TELECOM_ALARMS)} ALARMAS Y PROTECCIONES OPERATIVAS • 15 SITIOS TELECOM"
    ws1["A1"].font = font_title
    ws1["A1"].fill = fill_dark
    ws1["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws1.row_dimensions[1].height = 36

    ws1.merge_cells("A2:K2")
    ws1["A2"] = "Ecosistema HyVision • Integración Nativa Device Profile 'HYBRICO_MICROGRID' en ThingsBoard & API HyNet (Stellar by New Sun Road)"
    ws1["A2"].font = font_subtitle
    ws1["A2"].fill = fill_dark
    ws1["A2"].alignment = Alignment(horizontal="center", vertical="center")
    ws1.row_dimensions[2].height = 22

    # Encabezados
    headers_alm = [
        "Código Alarma",
        "Categoría / Subsistema",
        "Equipo / Elemento Físico",
        "Variable de Telemetría",
        "Condición de Disparo (Trigger)",
        "Condición de Despeje (Clear)",
        "Severidad",
        "Nivel de Implementación / Capa",
        "Descripción Técnica del Evento",
        "Protocolo Operativo / Acción Recomendada",
        "Sitios donde Aplica"
    ]

    ws1.row_dimensions[4].height = 30
    for c_idx, h in enumerate(headers_alm, 1):
        cell = ws1.cell(row=4, column=c_idx, value=h)
        cell.font = font_head
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border_thin

    for idx, a in enumerate(TELECOM_ALARMS, start=5):
        ws1.row_dimensions[idx].height = 38
        is_crit = a["severidad"].lower() == "critica"
        row_fill = fill_crit if is_crit else (fill_warn if a["severidad"].lower() == "advertencia" else (fill_zebra if idx % 2 == 1 else fill_white))

        row_vals = [
            a["id"],
            a["subsistema"],
            a["equipo"],
            a["variable"],
            a["disparo"],
            a["despeje"],
            a["severidad"].upper(),
            a["origen"],
            a["descripcion"],
            a["accion"],
            a.get("sitios_aplica", "Todos los 15 sitios Telecom")
        ]

        for c_idx, val in enumerate(row_vals, 1):
            cell = ws1.cell(row=idx, column=c_idx, value=val)
            cell.border = border_thin
            c_fill = row_fill if (c_idx == 7) else (fill_zebra if idx % 2 == 1 else fill_white)
            cell.fill = c_fill

            if c_idx == 1:
                cell.font = font_code
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif c_idx == 7:
                cell.font = font_crit if is_crit else font_warn
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif c_idx in [5, 6]:
                cell.font = font_cond
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            elif c_idx in [2, 3]:
                cell.font = font_bold
                cell.alignment = Alignment(vertical="center", wrap_text=True)
            elif c_idx == 4:
                cell.font = font_code
                cell.alignment = Alignment(vertical="center", wrap_text=True)
            else:
                cell.font = font_body
                cell.alignment = Alignment(vertical="center", wrap_text=True)

    col_widths_alm = [16, 24, 26, 28, 22, 22, 16, 32, 48, 52, 34]
    for i, w in enumerate(col_widths_alm, start=1):
        ws1.column_dimensions[get_column_letter(i)].width = w
    ws1.freeze_panes = "A5"
    ws1.auto_filter.ref = f"A4:{get_column_letter(len(headers_alm))}{len(TELECOM_ALARMS) + 4}"

    # =========================================================================
    # HOJA 2: Topologia_15_Sitios_Telecom
    # =========================================================================
    ws2 = wb.create_sheet(title="Topologia_15_Sitios_Telecom")
    ws2.views.sheetView[0].showGridLines = True

    ws2.merge_cells("A1:G1")
    ws2["A1"] = "CLASIFICACIÓN TOPOLÓGICA Y EQUIPAMIENTO DE LOS 15 SITIOS TELECOM"
    ws2["A1"].font = font_title
    ws2["A1"].fill = fill_dark
    ws2["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws2.row_dimensions[1].height = 34

    headers_top = [
        "Radiobase / Alias Sitio",
        "País",
        "ID Dispositivo ThingsBoard / Gateway (ps_id)",
        "Clasificación Topológica",
        "Equipamiento Solar Fotovoltaico",
        "Acometida Red Comercial",
        "Grupo Electrógeno Diésel"
    ]

    ws2.row_dimensions[3].height = 28
    for c_idx, h in enumerate(headers_top, 1):
        cell = ws2.cell(row=3, column=c_idx, value=h)
        cell.font = font_head
        cell.fill = fill_sub
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border_thin

    for idx, s in enumerate(SITES_15_LIST, start=4):
        ws2.row_dimensions[idx].height = 26
        row_vals = [
            s["alias"],
            s["pais"],
            s["ps_id"],
            s["topologia"],
            s["solar"],
            s["red"],
            s["dg"]
        ]
        for c_idx, val in enumerate(row_vals, 1):
            cell = ws2.cell(row=idx, column=c_idx, value=val)
            cell.border = border_thin
            cell.fill = fill_zebra if idx % 2 == 1 else fill_white
            if c_idx == 1:
                cell.font = font_bold
                cell.alignment = Alignment(vertical="center")
            elif c_idx == 3:
                cell.font = font_code
                cell.alignment = Alignment(vertical="center")
            elif c_idx in [2, 6, 7]:
                cell.font = font_body
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.font = font_body
                cell.alignment = Alignment(vertical="center", wrap_text=True)

    col_widths_top = [34, 16, 36, 26, 36, 20, 20]
    for i, w in enumerate(col_widths_top, start=1):
        ws2.column_dimensions[get_column_letter(i)].width = w
    ws2.freeze_panes = "A4"
    ws2.auto_filter.ref = f"A3:{get_column_letter(len(headers_top))}{len(SITES_15_LIST) + 3}"

    # =========================================================================
    # HOJA 3: Resumen_por_Subsistema
    # =========================================================================
    ws3 = wb.create_sheet(title="Resumen_por_Subsistema")
    ws3.views.sheetView[0].showGridLines = True

    ws3.merge_cells("A1:E1")
    ws3["A1"] = "RESUMEN DE PROTECCIONES POR SUBSISTEMA TELECOM"
    ws3["A1"].font = font_title
    ws3["A1"].fill = fill_dark
    ws3["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws3.row_dimensions[1].height = 34

    headers_res = ["Subsistema", "Total Alarmas", "Críticas (Disparo)", "Advertencias (Alerta)", "% Críticas"]
    ws3.row_dimensions[3].height = 26
    for c_idx, h in enumerate(headers_res, 1):
        cell = ws3.cell(row=3, column=c_idx, value=h)
        cell.font = font_head
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border_thin

    subsystems_summary = {}
    for a in TELECOM_ALARMS:
        sub = a["subsistema"]
        if sub not in subsystems_summary:
            subsystems_summary[sub] = {"total": 0, "crit": 0, "warn": 0}
        subsystems_summary[sub]["total"] += 1
        if a["severidad"].lower() == "critica":
            subsystems_summary[sub]["crit"] += 1
        else:
            subsystems_summary[sub]["warn"] += 1

    tot_all = 0
    tot_crit = 0
    tot_warn = 0
    for idx, (sub, data) in enumerate(subsystems_summary.items(), start=4):
        ws3.row_dimensions[idx].height = 24
        tot_all += data["total"]
        tot_crit += data["crit"]
        tot_warn += data["warn"]
        pct_crit = f"{(data['crit'] / data['total']) * 100:.1f}%"

        vals = [sub, data["total"], data["crit"], data["warn"], pct_crit]
        for c_idx, val in enumerate(vals, 1):
            cell = ws3.cell(row=idx, column=c_idx, value=val)
            cell.border = border_thin
            cell.fill = fill_zebra if idx % 2 == 1 else fill_white
            if c_idx == 1:
                cell.font = font_bold
                cell.alignment = Alignment(vertical="center")
            elif c_idx == 3:
                cell.font = font_crit
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif c_idx == 4:
                cell.font = font_warn
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.font = font_body
                cell.alignment = Alignment(horizontal="center", vertical="center")

    # Fila de Totales
    row_tot = len(subsystems_summary) + 4
    ws3.row_dimensions[row_tot].height = 28
    tot_vals = ["TOTAL SISTEMA TELECOM", tot_all, tot_crit, tot_warn, f"{(tot_crit / tot_all) * 100:.1f}%"]
    for c_idx, val in enumerate(tot_vals, 1):
        cell = ws3.cell(row=row_tot, column=c_idx, value=val)
        cell.font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
        cell.fill = fill_dark
        cell.border = border_thin
        cell.alignment = Alignment(horizontal="center", vertical="center")

    col_widths_res = [32, 18, 22, 24, 18]
    for i, w in enumerate(col_widths_res, start=1):
        ws3.column_dimensions[get_column_letter(i)].width = w
    ws3.freeze_panes = "A4"

    # Eliminar hoja default vacía
    if default_sheet in wb.worksheets:
        wb.remove(default_sheet)

    wb.save(NEW_EXCEL_PATH)
    print(f"-> EXCEL EXCLUSIVO TELECOM CREADO CON ÉXITO: {NEW_EXCEL_PATH}")

def restaurar_hoja_telecom_en_gaori():
    print(f"[2/3] Restaurando pestaña Telecom en el Excel de GAORI: {GAORI_EXCEL_PATH}...")
    try:
        wb = openpyxl.load_workbook(GAORI_EXCEL_PATH)
        sheet_name = "Alarmas_Telecom_15_Sitios"
        if sheet_name in wb.sheetnames:
            wb.remove(wb[sheet_name])

        ws = wb.create_sheet(title=sheet_name)
        ws.views.sheetView[0].showGridLines = True

        fill_dark    = PatternFill(start_color="0B1A10", end_color="0B1A10", fill_type="solid")
        fill_header  = PatternFill(start_color="1B4332", end_color="1B4332", fill_type="solid")
        fill_zebra   = PatternFill(start_color="F4FBF7", end_color="F4FBF7", fill_type="solid")
        fill_white   = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
        fill_crit    = PatternFill(start_color="FDE8E8", end_color="FDE8E8", fill_type="solid")
        fill_warn    = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")

        font_title   = Font(name="Segoe UI", size=15, bold=True, color="FFFFFF")
        font_subtitle= Font(name="Segoe UI", size=10, italic=True, color="E2E8F0")
        font_head    = Font(name="Segoe UI", size=10.5, bold=True, color="FFFFFF")
        font_body    = Font(name="Segoe UI", size=9.5, color="1E293B")
        font_bold    = Font(name="Segoe UI", size=9.5, bold=True, color="0F172A")
        font_code    = Font(name="Consolas", size=9.5, bold=True, color="0369A1")
        font_cond    = Font(name="Segoe UI", size=9.5, color="1E3A8A")
        font_crit    = Font(name="Segoe UI", size=9.5, bold=True, color="991B1B")
        font_warn    = Font(name="Segoe UI", size=9.5, bold=True, color="92400E")

        border_thin  = Border(
            left=Side(style='thin', color='CBD5E1'),
            right=Side(style='thin', color='CBD5E1'),
            top=Side(style='thin', color='CBD5E1'),
            bottom=Side(style='thin', color='CBD5E1')
        )

        ws.merge_cells("A1:K1")
        ws["A1"] = f"MATRIZ SCADA DE {len(TELECOM_ALARMS)} ALARMAS Y PROTECCIONES OPERATIVAS • 15 SITIOS TELECOM"
        ws["A1"].font = font_title
        ws["A1"].fill = fill_dark
        ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 36

        ws.merge_cells("A2:K2")
        ws["A2"] = "Ecosistema HyVision • Integración Nativa Device Profile 'HYBRICO_MICROGRID' en ThingsBoard & API HyNet"
        ws["A2"].font = font_subtitle
        ws["A2"].fill = fill_dark
        ws["A2"].alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[2].height = 22

        headers_alm = [
            "Código Alarma", "Categoría / Subsistema", "Equipo / Elemento Físico",
            "Variable de Telemetría", "Condición de Disparo (Trigger)", "Condición de Despeje (Clear)",
            "Severidad", "Nivel de Implementación / Capa", "Descripción Técnica del Evento",
            "Protocolo Operativo / Acción Recomendada", "Sitios donde Aplica"
        ]

        ws.row_dimensions[4].height = 30
        for c_idx, h in enumerate(headers_alm, 1):
            cell = ws.cell(row=4, column=c_idx, value=h)
            cell.font = font_head
            cell.fill = fill_header
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = border_thin

        for idx, a in enumerate(TELECOM_ALARMS, start=5):
            ws.row_dimensions[idx].height = 38
            is_crit = a["severidad"].lower() == "critica"
            row_fill = fill_crit if is_crit else (fill_warn if a["severidad"].lower() == "advertencia" else (fill_zebra if idx % 2 == 1 else fill_white))

            row_vals = [
                a["id"], a["subsistema"], a["equipo"], a["variable"],
                a["disparo"], a["despeje"], a["severidad"].upper(),
                a["origen"], a["descripcion"], a["accion"],
                a.get("sitios_aplica", "Todos los 15 sitios Telecom")
            ]

            for c_idx, val in enumerate(row_vals, 1):
                cell = ws.cell(row=idx, column=c_idx, value=val)
                cell.border = border_thin
                c_fill = row_fill if (c_idx == 7) else (fill_zebra if idx % 2 == 1 else fill_white)
                cell.fill = c_fill

                if c_idx == 1:
                    cell.font = font_code
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                elif c_idx == 7:
                    cell.font = font_crit if is_crit else font_warn
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                elif c_idx in [5, 6]:
                    cell.font = font_cond
                    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                elif c_idx in [2, 3]:
                    cell.font = font_bold
                    cell.alignment = Alignment(vertical="center", wrap_text=True)
                elif c_idx == 4:
                    cell.font = font_code
                    cell.alignment = Alignment(vertical="center", wrap_text=True)
                else:
                    cell.font = font_body
                    cell.alignment = Alignment(vertical="center", wrap_text=True)

        col_widths_alm = [16, 24, 26, 28, 22, 22, 16, 32, 48, 52, 34]
        for i, w in enumerate(col_widths_alm, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A5"
        ws.auto_filter.ref = f"A4:{get_column_letter(len(headers_alm))}{len(TELECOM_ALARMS) + 4}"

        wb.save(GAORI_EXCEL_PATH)
        print("-> Pestaña 'Alarmas_Telecom_15_Sitios' restaurada exitosamente en el Excel de GAORI.")
    except Exception as e:
        print(f"Nota sobre Excel de GAORI (puede estar abierto en Excel): {e}")

if __name__ == "__main__":
    generar_excel_dedicado_telecom()
    restaurar_hoja_telecom_en_gaori()
